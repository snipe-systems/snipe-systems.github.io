"""Pluggable journal sinks: CSV, JSON Lines, and (optional) XLSX.

A `Journal` writes to one or more sinks at once, so every trade can go to your
plain CSV *and* a JSON log *and* an Excel sheet simultaneously — no changes to
the trading logic required.
"""

from __future__ import annotations

import csv
import json
import time
from abc import ABC, abstractmethod
from pathlib import Path

COLUMNS = ["timestamp", "symbol", "action", "quantity", "fill_price",
           "strategy", "idempotency_key"]

# Characters that make spreadsheet apps treat a cell as a formula.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _sanitize(value) -> str:
    """Neutralize spreadsheet formula injection in user-controlled cells."""
    s = "" if value is None else str(value)
    if s.startswith(_FORMULA_PREFIXES):
        return "'" + s
    return s


class JournalSink(ABC):
    @abstractmethod
    def write(self, row: dict) -> None: ...

    @abstractmethod
    def read(self) -> list[dict]: ...

    def tag_last(self, tag: str) -> bool:
        """Apply a strategy tag to the most recent row. Returns False if empty."""
        rows = self.read()
        if not rows:
            return False
        rows[-1]["strategy"] = tag
        self._rewrite(rows)
        return True

    def _rewrite(self, rows: list[dict]) -> None:  # pragma: no cover - overridden
        raise NotImplementedError


class CsvSink(JournalSink):
    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if not self.path.exists() or self.path.stat().st_size == 0:
            with open(self.path, "a", newline="") as f:
                csv.writer(f).writerow(COLUMNS)

    def write(self, row: dict) -> None:
        with open(self.path, "a", newline="") as f:
            csv.writer(f).writerow([_sanitize(row.get(c, "")) for c in COLUMNS])

    def read(self) -> list[dict]:
        if not self.path.exists():
            return []
        with open(self.path, newline="") as f:
            return [dict(r) for r in csv.DictReader(f)]

    def _rewrite(self, rows: list[dict]) -> None:
        with open(self.path, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=COLUMNS)
            w.writeheader()
            for r in rows:
                w.writerow({c: r.get(c, "") for c in COLUMNS})


class JsonLinesSink(JournalSink):
    """One JSON object per line (append-friendly, easy for other tools to read)."""

    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)

    def write(self, row: dict) -> None:
        with open(self.path, "a") as f:
            f.write(json.dumps(row) + "\n")

    def read(self) -> list[dict]:
        if not self.path.exists():
            return []
        rows = []
        with open(self.path) as f:
            for line in f:
                line = line.strip()
                if line:
                    try:
                        rows.append(json.loads(line))
                    except json.JSONDecodeError:
                        continue
        return rows

    def _rewrite(self, rows: list[dict]) -> None:
        with open(self.path, "w") as f:
            for r in rows:
                f.write(json.dumps(r) + "\n")


class XlsxSink(JournalSink):
    """Excel export. Requires `openpyxl` (pip install openpyxl)."""

    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._rows: list[dict] = self.read()

    def write(self, row: dict) -> None:
        self._rows.append(row)
        self._write_xlsx()

    def read(self) -> list[dict]:
        if not self.path.exists():
            return []
        try:
            from openpyxl import load_workbook
        except ImportError:
            return []
        wb = load_workbook(self.path)
        ws = wb.active
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                header = list(row)
                continue
            rows.append(dict(zip(header, row)))
        return rows

    def _rewrite(self, rows: list[dict]) -> None:
        self._rows = rows
        self._write_xlsx()

    def _write_xlsx(self) -> None:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(COLUMNS)
        for r in self._rows:
            ws.append([r.get(c, "") for c in COLUMNS])
        wb.save(self.path)


def build_sink(path: str | Path, fmt: str = "csv") -> JournalSink:
    fmt = (fmt or "csv").lower()
    if fmt in ("csv",):
        return CsvSink(path)
    if fmt in ("json", "jsonl", "ndjson"):
        return JsonLinesSink(path)
    if fmt in ("xlsx", "excel"):
        return XlsxSink(path)
    raise ValueError(f"unknown journal format: {fmt}")
